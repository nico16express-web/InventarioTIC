"""Inventario Tecnológico Escolar - versión web (Flask + SQLite)."""
import hmac, io, os, sqlite3
from datetime import datetime
from flask import Flask, Response, jsonify, request, render_template, send_file, g
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

DB = os.environ.get("INVENTARIO_DB", os.path.join(os.path.dirname(os.path.abspath(__file__)), "inventario_escolar.db"))
DATABASE_URL = os.environ.get("DATABASE_URL")  # si existe -> Postgres (Vercel); si no -> SQLite local
PG = bool(DATABASE_URL)
if PG:
    import psycopg
    from psycopg.rows import dict_row
    IntegrityError = psycopg.errors.UniqueViolation
else:
    IntegrityError = sqlite3.IntegrityError

app = Flask(__name__, static_folder="public", static_url_path="")

class Conn:
    def __init__(self):
        if PG:
            self.c = psycopg.connect(DATABASE_URL, row_factory=dict_row)
        else:
            self.c = sqlite3.connect(DB); self.c.row_factory = sqlite3.Row
    def execute(self, sql, p=()):
        return self.c.execute(sql.replace("?", "%s") if PG else sql, p)
    def commit(self): self.c.commit()
    def close(self): self.c.close()

@app.before_request
def acceso():
    u, p = os.environ.get("APP_USER"), os.environ.get("APP_PASS")
    if u and p:
        a = request.authorization
        if not (a and hmac.compare_digest(a.username or "", u) and hmac.compare_digest(a.password or "", p)):
            return Response("Acceso restringido", 401, {"WWW-Authenticate": 'Basic realm="Inventario"'})

LISTAS = {
    "niveles": ["Básica", "Parvularia", "No Aplica"],
    "financiamiento": ["Inversión", "Donación", "Sin información"],
    "adquisicion": ["Subvencion Regular", "Fondos SEP", "Sin Información"],
    "estados": ["En uso", "Nuevo", "Descartable"],
    "categorias": [
        "Implementos deportivos", "Implementos de laboratorio",
        "Instrumentos musicales y/o artísticos", "Libros y revistas", "Equipos informáticos",
        "Equipos multicopiadores", "Equipos de amplificación y sonido",
        "Equipos de climatización: calefacción, ventilación y aire acondicionado",
        "Útiles escolares, equipamiento especializado de Liceos Técnicos Profesionales o materiales de oficina",
        "Materiales y útiles de aseo", "Mobiliario escolar fuera de la sala de clase",
        "Mobiliario escolar dentro de la sala de clase", "Mobiliario no pedagógico o de oficina",
        "Otros equipos, materiales o insumos"],
}
CAMPOS = ["id", "nivel", "objeto_tecnologico", "cantidad", "dependencia", "financiamiento", "adquisicion", "categoria", "estado", "fecha"]

def db():
    if "db" not in g:
        g.db = Conn()
    return g.db

@app.teardown_appcontext
def cerrar(_):
    c = g.pop("db", None)
    if c: c.close()

def init_db():
    c = Conn()
    c.execute("""CREATE TABLE IF NOT EXISTS inventario (id TEXT PRIMARY KEY, nivel TEXT, objeto_tecnologico TEXT,
        cantidad INTEGER, dependencia TEXT, financiamiento TEXT, adquisicion TEXT, categoria TEXT, estado TEXT, fecha TEXT)""")
    pk = "SERIAL PRIMARY KEY" if PG else "INTEGER PRIMARY KEY AUTOINCREMENT"
    c.execute(f"""CREATE TABLE IF NOT EXISTS movimientos (n {pk}, item_id TEXT,
        objeto TEXT, accion TEXT, detalle TEXT, fecha TEXT)""")
    c.commit(); c.close()

def log(item_id, objeto, accion, detalle=""):
    db().execute("INSERT INTO movimientos(item_id,objeto,accion,detalle,fecha) VALUES(?,?,?,?,?)",
                 (item_id, objeto, accion, detalle, datetime.now().strftime("%d/%m/%Y %H:%M")))

def validar(d, nuevo):
    d = {k: str(d.get(k) or "").strip() for k in CAMPOS}
    if nuevo and not d["id"]: raise ValueError("Falta el ID / Código.")
    if not d["objeto_tecnologico"]: raise ValueError("Falta el nombre del objeto.")
    try:
        d["cantidad"] = int(d["cantidad"] or 0)
        if d["cantidad"] < 0: raise ValueError
    except ValueError:
        raise ValueError("La cantidad debe ser un número entero, 0 o mayor.")
    d["fecha"] = d["fecha"] or "Sin información"
    return d

def filtrar(a):
    sql, p = "SELECT * FROM inventario WHERE 1=1", []
    if a.get("q"):
        sql += " AND (LOWER(id) LIKE ? OR LOWER(objeto_tecnologico) LIKE ?)"; p += [f"%{a['q'].lower()}%"] * 2
    for campo in ("categoria", "estado", "dependencia"):
        if a.get(campo):
            sql += f" AND {campo}=?"; p.append(a[campo])
    return db().execute(sql + " ORDER BY objeto_tecnologico", p).fetchall()

@app.route("/")
def index():
    return render_template("index.html", listas=LISTAS)

@app.get("/api/items")
def listar():
    return jsonify([dict(r) for r in filtrar(request.args)])

@app.post("/api/items")
def crear():
    try:
        d = validar(request.get_json(force=True), True)
        db().execute(f"INSERT INTO inventario VALUES ({','.join('?' * len(CAMPOS))})", [d[k] for k in CAMPOS])
        log(d["id"], d["objeto_tecnologico"], "Alta", f"Stock inicial: {d['cantidad']}")
        db().commit()
        return jsonify(d), 201
    except ValueError as e:
        return jsonify(error=str(e)), 400
    except IntegrityError:
        return jsonify(error="Ese ID ya existe. Usa un código distinto."), 409

@app.put("/api/items/<path:item_id>")
def editar(item_id):
    try:
        d = validar({**request.get_json(force=True), "id": item_id}, False)
        cur = db().execute("""UPDATE inventario SET nivel=?, objeto_tecnologico=?, cantidad=?, dependencia=?,
            financiamiento=?, adquisicion=?, categoria=?, estado=?, fecha=? WHERE id=?""",
            [d[k] for k in CAMPOS[1:]] + [item_id])
        if not cur.rowcount: return jsonify(error="El registro no existe."), 404
        log(item_id, d["objeto_tecnologico"], "Edición", f"Stock: {d['cantidad']}, estado: {d['estado']}")
        db().commit()
        return jsonify(d)
    except ValueError as e:
        return jsonify(error=str(e)), 400

@app.delete("/api/items/<path:item_id>")
def eliminar(item_id):
    r = db().execute("SELECT objeto_tecnologico FROM inventario WHERE id=?", (item_id,)).fetchone()
    if not r: return jsonify(error="El registro no existe."), 404
    db().execute("DELETE FROM inventario WHERE id=?", (item_id,))
    log(item_id, r["objeto_tecnologico"], "Baja", "Registro eliminado")
    db().commit()
    return jsonify(ok=True)

@app.post("/api/descontar/<path:item_id>")
def descontar(item_id):
    try:
        n = int((request.get_json(force=True) or {}).get("cantidad", 0))
    except (TypeError, ValueError):
        return jsonify(error="Ingresa un número válido."), 400
    if n <= 0: return jsonify(error="La cantidad a descontar debe ser mayor a cero."), 400
    r = db().execute("SELECT cantidad, objeto_tecnologico FROM inventario WHERE id=?", (item_id,)).fetchone()
    if not r: return jsonify(error="El registro no existe."), 404
    if r["cantidad"] < n: return jsonify(error=f"No hay suficiente stock. Stock actual: {r['cantidad']}."), 400
    nuevo = r["cantidad"] - n
    db().execute("UPDATE inventario SET cantidad=? WHERE id=?", (nuevo, item_id))
    log(item_id, r["objeto_tecnologico"], "Descuento", f"-{n} (de {r['cantidad']} a {nuevo})")
    db().commit()
    return jsonify(cantidad=nuevo)

@app.get("/api/movimientos")
def movimientos():
    rows = db().execute("SELECT * FROM movimientos ORDER BY n DESC LIMIT 200").fetchall()
    return jsonify([dict(r) for r in rows])

@app.get("/api/export")
def exportar():
    rows = filtrar(request.args)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Inventario Tecnológico"
    ws.append(["Nivel", "ID / Código", "Objeto Tecnológico", "Stock Actual", "Dependencia", "Financiamiento",
               "Adquisición", "Categoría", "Estado", "Fecha"])
    for c in ws[1]:
        c.fill = PatternFill("solid", start_color="083668"); c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r[k] for k in ["nivel", "id", "objeto_tecnologico", "cantidad", "dependencia", "financiamiento",
                                  "adquisicion", "categoria", "estado", "fecha"]])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(len(str(c.value or "")) for c in col) + 2)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Inventario_Escuela_{datetime.now():%Y%m%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

init_db()
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)
